#!/usr/bin/env python3
"""Build the journal-style DOCX manuscript from the project draft and results."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
TEX = ROOT / "paper" / "paper_draft.tex"
NUMBERS = ROOT / "paper" / "numbers.tex"
BIB = ROOT / "paper" / "references.bib"
FIGURES = ROOT / "paper" / "figures_submission"
EQUATIONS = ROOT / "paper" / "equations"
OUTPUT = ROOT / "output" / "manuscript"
OUTPUT.mkdir(parents=True, exist_ok=True)

TITLE = "Compulsions as epistemic traps: repetition obscures causal control"
AUTHOR = "Dai Miyazaki"
AFFILIATION = "Graduate School of Humanities and Sociology, The University of Tokyo"
CORRESPONDENCE = "dmyzk29@gmail.com"
ORCID = "https://orcid.org/0009-0000-0829-4435"
REPOSITORY = "https://github.com/buchendschmerz/compulsions-as-epistemic-traps"

ABSTRACT = (
    "Compulsive safety behaviour is commonly explained by biased beliefs about danger or control, or "
    "by impaired learning. This article identifies a complementary mechanism that requires neither "
    "deficit. In a hazard model, an episode starts in danger with probability q, and each preventive "
    "repetition succeeds with probability p. If an agent always repeats the action k times, outcomes "
    "reveal only q(1-p)ᵏ. The Fisher information is rank one, so q and p are structurally "
    "non-identifiable under any fixed-k policy, regardless of sample size. This theorem does not "
    "depend on a biased prior. Repetition removes the variation needed to separate the two causes. "
    "Simulations of exact Bayesian learning show why this geometry matters: with an initially "
    "pessimistic prior, a repetitive agent becomes precise about the observed harm probability while "
    "remaining wrong about its causes; varying k recovers both parameters. Interventions restore "
    "information only when they create contrasts between action counts. Uncertainty aversion increases "
    "repetition and worsens causal learning, but it does not sustain repetition over the simulated "
    "horizon. The model therefore explains why causal beliefs can remain wrong; it does not, by itself, "
    "explain why repetitive behaviour persists. It predicts a direct human test: schedules matched on "
    "mean action count but differing in where they vary should produce different learning about personal "
    "efficacy, with expected information outperforming raw behavioural variance. The result is a "
    "falsifiable mechanism for failed causal learning, not a complete account of obsessive-compulsive "
    "disorder."
)

KEYWORDS = (
    "obsessive-compulsive disorder; structural identifiability; Bayesian learning; safety behaviour; "
    "exposure and response prevention"
)

INK = "18242F"
BLUE = "2878B5"
MUTED = "68727D"
LIGHT = "EEF3F6"
PALE_BLUE = "E5F0F7"
PALE_GOLD = "FFF5D8"
WHITE = "FFFFFF"


def set_run_font(run, name="Arial", size=None, bold=None, italic=None, color=INK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=8.5, color=MUTED)


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.82)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.34)
    section.footer_distance = Inches(0.34)
    doc.settings.odd_and_even_pages_header_footer = False
    section.different_first_page_header_footer = False

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.17
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    heading_tokens = {
        "Heading 1": (15.0, 18, 7, BLUE),
        "Heading 2": (12.0, 12, 5, INK),
        "Heading 3": (10.8, 9, 3, INK),
    }
    for name, (size, before, after, color) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "Arial"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    caption.font.size = Pt(8.5)
    caption.font.bold = False
    caption.font.color.rgb = RGBColor.from_string(INK)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(9)
    caption.paragraph_format.line_spacing = 1.05
    caption.paragraph_format.keep_together = True

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    set_run_font(p.add_run("Page "), size=8.5, color=MUTED)
    add_page_field(p)


def parse_macros(text):
    return {name: value for name, value in re.findall(
        r"\\newcommand\{\\([A-Za-z]+)\}\{([^{}]*(?:\{[^{}]*\}[^{}]*)*)\}", text
    )}


def parse_bib(text):
    starts = list(re.finditer(r"(?m)^@(\w+)\{([^,]+),", text))
    entries = {}
    field_re = re.compile(r"(?ms)^\s*([A-Za-z][A-Za-z0-9_-]*)\s*=\s*\{(.*?)\}\s*,?\s*$")
    for index, match in enumerate(starts):
        end = starts[index + 1].start() if index + 1 < len(starts) else len(text)
        block = text[match.start():end]
        fields = {key.lower(): " ".join(value.split()) for key, value in field_re.findall(block)}
        entries[match.group(2).strip()] = fields
    return entries


def author_year(key, entries, narrative=False):
    entry = entries.get(key, {})
    authors = entry.get("author", key).split(" and ")
    authors = [author for author in authors if author.lower() != "others"]
    surnames = [decode_latex_text(author.split(",", 1)[0].strip()) for author in authors]
    if not surnames:
        lead = key
    elif len(surnames) == 1:
        lead = surnames[0]
    elif len(surnames) == 2:
        lead = f"{surnames[0]} and {surnames[1]}"
    else:
        lead = f"{surnames[0]} et al."
    year = entry.get("year", "n.d.")
    return f"{lead} ({year})" if narrative else f"{lead}, {year}"


REF_MAP = {
    "fig:fig1": "1", "fig:fig2": "2", "fig:fig3": "3", "fig:fig6": "4",
    "fig:fig7": "5", "fig:fig4": "S1", "fig:fig5": "S2",
    "fig:figS2": "S3", "tab:accounts": "2", "prop:rank1": "1",
    "tab:params": "1", "prop:two": "2", "cor:ridge": "1",
    "sec:theorem": "2.3", "sec:baseline": "2.1", "sec:interventions": "3.5",
    "eq:lambda": "(1)", "eq:cost": "(2)", "eq:fisher": "(3)",
    "eq:slope": "(4)", "eq:recovery": "(5)",
}


SUBSCRIPT_TRANSLATION = str.maketrans(
    "0123456789+-=()aehijklmnoprstx",
    "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎ₐₑₕᵢⱼₖₗₘₙₒₚᵣₛₜₓ",
)
SUPERSCRIPT_TRANSLATION = str.maketrans(
    "0123456789+-=()abcdefghijklmnoprstuvwxyzT",
    "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ᵃᵇᶜᵈᵉᶠᵍʰⁱʲᵏˡᵐⁿᵒᵖʳˢᵗᵘᵛʷˣʸᶻᵀ",
)


def prettify_inline_math(text):
    def script(group, translation, marker):
        compact = group.replace("_", "").replace(" ", "")
        converted = compact.translate(translation)
        if all(ord(char) in translation for char in compact):
            return converted
        return marker + group.replace("_", "")

    text = text.replace("^*", "*")
    text = re.sub(r"_\(([^()]*)\)",
                  lambda m: script(m.group(1), SUBSCRIPT_TRANSLATION, "_"), text)
    text = re.sub(r"_([A-Za-z0-9+\-=]+)",
                  lambda m: script(m.group(1), SUBSCRIPT_TRANSLATION, "_"), text)
    text = re.sub(r"\^\(([^()]*)\)",
                  lambda m: script(m.group(1), SUPERSCRIPT_TRANSLATION, "^"), text)
    text = re.sub(r"\^([A-Za-z0-9+\-=]+)",
                  lambda m: script(m.group(1), SUPERSCRIPT_TRANSLATION, "^"), text)
    return text


def latex_plain(text, entries):
    text = re.sub(r"\\citet\{([^}]+)\}",
                  lambda m: "; ".join(author_year(k.strip(), entries, True) for k in m.group(1).split(",")), text)
    text = re.sub(r"\\citep\{([^}]+)\}",
                  lambda m: "(" + "; ".join(author_year(k.strip(), entries) for k in m.group(1).split(",")) + ")", text)
    text = re.sub(r"\\ref\{([^}]+)\}", lambda m: REF_MAP.get(m.group(1), m.group(1)), text)
    text = re.sub(r"\\eqref\{([^}]+)\}", lambda m: REF_MAP.get(m.group(1), m.group(1)), text)
    text = re.sub(r"\$([^$]+)\$", "\ue000\\1\ue001", text)
    replacements = {
        r"\\mathbb{E}": "E", r"\\mathcal{I}": "I", r"\\Ifisher": "I",
        r"\\lamk": "λ_k", r"\\lambda": "λ", r"\\theta": "θ", r"\\pi": "π",
        r"\\alpha": "α", r"\\kappa": "κ", r"\\tau": "τ", r"\\rho": "ρ",
        r"\\Delta": "Δ", r"\\approx": "≈", r"\\neq": "≠", r"\\ge": "≥",
        r"\\le": "≤", r"\\in": "∈", r"\\times": "×", r"\\pm": "±",
        r"\\mid": "|", r"\\ast": "*", r"\\top": "T", r"\\nats": "nats",
        r"\\sim": "∼", r"\\nabla": "∇", r"\\det": "det", r"\\rightarrow": "→",
        r"\\Longrightarrow": "⇒", r"\\sum": "Σ", r"\\log": "log",
        r"\\to": "→", r"\\dots": "...",
    }
    for source, target in replacements.items():
        text = re.sub(source, target, text)
    text = text.replace(r"\{", "<<LBRACE>>").replace(r"\}", "<<RBRACE>>")
    text = text.replace(r"\%", "%").replace(r"\_", "_")
    text = re.sub(r"\\(?:emph|textbf|texttt|mathrm|text|operatorname)\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\frac\{([^{}]+)\}\{([^{}]+)\}", r"(\1)/(\2)", text)
    text = re.sub(r"\\begin\{[^}]+\}|\\end\{[^}]+\}", "", text)
    text = re.sub(r"\\label\{[^}]+\}", "", text)
    text = text.replace("\\,", " ").replace("\\;", " ").replace("\\!", "")
    text = text.replace(r"\left", "").replace(r"\right", "").replace(r"\bigg", "")
    text = text.replace(r"\qquad", "   ").replace(r"\quad", "  ")
    text = re.sub(r"\\\\(?:\[[^]]+\])?", "; ", text)
    text = text.replace("~", " ").replace("$", "")
    text = text.replace("``", '"').replace("''", '"')
    text = text.replace("---", " - ").replace("--", "-")
    text = re.sub(r"\^\{([^{}]+)\}", r"^(\1)", text)
    text = re.sub(r"_\{([^{}]+)\}", r"_(\1)", text)
    text = text.replace("{", "").replace("}", "")
    text = text.replace("<<LBRACE>>", "{").replace("<<RBRACE>>", "}")
    text = re.sub(r"\\[A-Za-z]+", "", text)
    text = re.sub("\ue000(.*?)\ue001", lambda m: prettify_inline_math(m.group(1)), text)
    text = re.sub(r"\s+", " ", text).strip()
    panel_updates = {
        "Figure 2D": "Figure 2B", "Figure 2E": "Figure 2C",
        "Fig. 2D": "Fig. 2B", "Fig. 2E": "Fig. 2C",
        "Figure 3C": "the learning summary", "Fig. 3C": "the learning summary",
        "Figure 3D": "Figure 3C", "Fig. 3D": "Fig. 3C",
        "Figure 3E": "Figure 3D", "Fig. 3E": "Fig. 3D",
        "Figure 3F": "the fixed-design sweep", "Fig. 3F": "the fixed-design sweep",
        "Figure 4D": "Figure 4C", "Fig. 4D": "Fig. 4C",
        "Figure 4E": "Figure 4B", "Fig. 4E": "Fig. 4B",
        "Figure 4F": "Figure 4D", "Fig. 4F": "Fig. 4D",
    }
    for source, target in panel_updates.items():
        text = text.replace(source, target)
    text = re.sub(r"\bFigure (?=[0-9S])", "Fig ", text)
    text = text.replace("Fig. ", "Fig ")
    text = text.replace("log_(10)", "log10")
    text = text.replace("pre-registered", "pre-specified")
    return text


EQUATION_ALTS = {
    1: "Equation 1. Harm is Bernoulli with probability lambda sub k equal to q times one minus p to the power k.",
    2: "Equation 2. Expected cost combines expected harm, uncertainty aversion and repetition cost; utility adds expected information gain.",
    3: "Equation 3. Fisher information is the outer product of the likelihood gradient divided by Bernoulli variance; the gradient is shown as a two-element column vector.",
    4: "Equation 4. Along a constant-likelihood ridge, d q over d p equals k q over one minus p and is non-negative.",
    5: "Equation 5. Two distinct repetition counts recover preventive efficacy p and baseline hazard q from their two outcome probabilities.",
}


def equation_number(text):
    labels = {"eq:lambda": 1, "eq:cost": 2, "eq:fisher": 3,
              "eq:slope": 4, "eq:recovery": 5}
    for label, number in labels.items():
        if label in text:
            return number
    raise ValueError(f"Unknown display equation: {text[:80]}")


def prepare_blocks(tex, macros):
    body = tex[tex.index(r"\section{Introduction}"):tex.index(r"\bibliography{references}")]
    for name, value in sorted(macros.items(), key=lambda item: -len(item[0])):
        body = re.sub(rf"\\{re.escape(name)}(?![A-Za-z])", lambda _: value, body)
    body = body.replace(r"\lamk", "λ_k").replace(r"\Ifisher", "I")
    body = re.sub(r"(?ms)\\begin\{figure\}.*?\\end\{figure\}", "\n\n", body)
    body = re.sub(r"(?ms)\\begin\{table\}.*?\\end\{table\}", "\n\n", body)
    body = re.sub(r"(?ms)\\section\*\{Data and code availability\}.*$", "", body)
    body = re.sub(r"(?m)^%.*$", "", body)
    body = body.replace(r"\itemsep2pt", "")

    stored = []

    def stash(kind, value):
        index = len(stored)
        stored.append((kind, value))
        return f"\n\n@@BLOCK_{index}@@\n\n"

    body = re.sub(r"(?ms)\\begin\{equation\}(.*?)\\end\{equation\}",
                  lambda m: stash("equation", m.group(1)), body)
    proposition_counter = {"proposition": 0, "corollary": 0}

    def theorem(match):
        kind = match.group(1)
        proposition_counter[kind] += 1
        title = (match.group(2) or "").strip()
        label = f"{kind.title()} {proposition_counter[kind]}"
        if title:
            label += f": {title}"
        theorem_text = re.sub(
            r"@@BLOCK_(\d+)@@",
            lambda nested: f"[[EQ:{equation_number(stored[int(nested.group(1))][1])}]]"
            if stored[int(nested.group(1))][0] == "equation" else "",
            match.group(3),
        )
        return stash("theorem", (label, theorem_text))

    body = re.sub(
        r"(?ms)\\begin\{(proposition|corollary)\}(?:\[([^]]+)\])?(.*?)\\end\{\1\}",
        theorem, body,
    )
    body = re.sub(r"(?ms)\\begin\{proof\}(.*?)\\end\{proof\}",
                  lambda m: stash("proof", m.group(1)), body)

    for command, kind in (("section*", "h1"), ("section", "h1"),
                          ("subsection", "h2"), ("subsubsection", "h3"),
                          ("paragraph", "phead")):
        body = re.sub(rf"(?ms)\\{re.escape(command)}\{{(.*?)\}}",
                      lambda m, kind=kind: stash(kind, m.group(1)), body)

    body = body.replace(r"\begin{enumerate}", "").replace(r"\end{enumerate}", "")
    body = re.sub(r"\\item\s*", lambda _: stash("item", ""), body)
    body = re.sub(r"\\label\{[^}]+\}", "", body)

    blocks = []
    for chunk in re.split(r"\n\s*\n", body):
        chunk = chunk.strip()
        if not chunk:
            continue
        marker = re.fullmatch(r"@@BLOCK_(\d+)@@", chunk)
        if marker:
            blocks.append(stored[int(marker.group(1))])
        else:
            blocks.append(("paragraph", " ".join(line.strip() for line in chunk.splitlines())))
    return blocks


def add_title_page(doc):
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_before = Pt(8)
    kicker.paragraph_format.space_after = Pt(9)
    set_run_font(kicker.add_run("RESEARCH ARTICLE"),
                 size=8.5, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(TITLE), size=23, bold=True, color=INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    set_run_font(p.add_run(AUTHOR), size=11.5, bold=True, color=INK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    set_run_font(p.add_run(AFFILIATION), size=9.5, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    set_run_font(p.add_run(f"ORCID: {ORCID}"), size=9.0, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    set_run_font(p.add_run(f"Correspondence: {CORRESPONDENCE}"), size=9.0, color=MUTED)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(8)
    p_pr = rule._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BLUE)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(3)
    h.paragraph_format.space_after = Pt(4)
    set_run_font(h.add_run("Abstract"), size=12.5, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    set_run_font(p.add_run(ABSTRACT), size=9.8, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    set_run_font(p.add_run("Keywords: "), size=9.0, bold=True, color=INK)
    set_run_font(p.add_run(KEYWORDS), size=9.0, color=MUTED)


def set_picture_alt(doc, title, description):
    shape = doc.inline_shapes[-1]
    shape._inline.docPr.set("title", title)
    shape._inline.docPr.set("descr", description)


def _math_run(text):
    run = OxmlElement("m:r")
    value = OxmlElement("m:t")
    value.set(qn("xml:space"), "preserve")
    value.text = str(text)
    run.append(value)
    return run


def _append_math(parent, value):
    if isinstance(value, (list, tuple)):
        for item in value:
            _append_math(parent, item)
    elif isinstance(value, str):
        parent.append(_math_run(value))
    else:
        parent.append(value)


def _math_script(tag, base, script):
    node = OxmlElement(f"m:{tag}")
    base_node = OxmlElement("m:e")
    _append_math(base_node, base)
    node.append(base_node)
    script_node = OxmlElement("m:sub" if tag == "sSub" else "m:sup")
    _append_math(script_node, script)
    node.append(script_node)
    return node


def _math_sub(base, subscript):
    return _math_script("sSub", base, subscript)


def _math_sup(base, superscript):
    return _math_script("sSup", base, superscript)


def _math_frac(numerator, denominator):
    node = OxmlElement("m:f")
    num = OxmlElement("m:num")
    den = OxmlElement("m:den")
    _append_math(num, numerator)
    _append_math(den, denominator)
    node.extend([num, den])
    return node


def _equation_lines(number):
    lam_k = lambda: _math_sub("λ", "k")
    lam_k1 = lambda: _math_sub("λ", "k₁")
    lam_k2 = lambda: _math_sub("λ", "k₂")
    grad_theta = lambda: _math_sub("∇", "θ")
    one_minus_p_k = lambda: _math_sup("(1 − p)", "k")
    if number == 1:
        return [[
            "Y ∼ Bernoulli(", lam_k(), "),    ", lam_k(), " = q", one_minus_p_k(),
        ]]
    if number == 2:
        return [
            ["C(k) = ", _math_sub("c", "harm"), "(E[", lam_k(), "] + κ SD[",
             lam_k(), "]) + ", _math_sub("c", "action"), "k"],
            ["U(k) = −C(k) + τ EIG(k)"],
        ]
    if number == 3:
        numerator = [grad_theta(), " ", lam_k(), " ",
                     _math_sup([grad_theta(), " ", lam_k()], "T")]
        denominator = [lam_k(), "(1 − ", lam_k(), ")"]
        gradient = _math_sup([
            "(", one_minus_p_k(), ",  −kq",
            _math_sup("(1 − p)", "k−1"), ")",
        ], "T")
        return [
            [_math_sub("I", "k"), "(θ) = ", _math_frac(numerator, denominator)],
            [grad_theta(), " ", lam_k(), " = ", gradient],
        ]
    if number == 4:
        ridge = _math_sub("│", [lam_k(), " constant"])
        return [[
            _math_frac("dq", "dp"), ridge, " = ", _math_frac("kq", "1 − p"), " ≥ 0",
        ]]
    if number == 5:
        exponent = _math_frac("1", "k₁ − k₂")
        return [
            [_math_frac(lam_k1(), lam_k2()), " = ",
             _math_sup("(1 − p)", "k₁−k₂")],
            ["p = 1 − ",
             _math_sup(["(", _math_frac(lam_k1(), lam_k2()), ")"], exponent),
             ",    q = ", lam_k1(), _math_sup("(1 − p)", "−k₁")],
        ]
    raise ValueError(f"Unknown equation number: {number}")


def add_equation_native(doc, number, paragraph=None):
    """Insert an editable Office Math equation rather than a raster image."""
    p = paragraph or doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_together = True
    p.paragraph_format.tab_stops.add_tab_stop(Inches(3.05), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.0), WD_TAB_ALIGNMENT.RIGHT)
    lines = _equation_lines(number)
    for index, components in enumerate(lines):
        if index:
            p.add_run().add_break()
        p.add_run("\t")
        equation = OxmlElement("m:oMath")
        _append_math(equation, components)
        p._p.append(equation)
        if index == len(lines) - 1:
            number_run = p.add_run(f"\t({number})")
            set_run_font(number_run, size=10.0, color=INK)
    return p


def add_figure(doc, stem, caption):
    image_path = FIGURES / f"{stem}.png"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image_path), width=Inches(6.42))
    set_picture_alt(doc, caption.split(". ", 1)[0], caption)
    c = doc.add_paragraph(style="Caption")
    set_run_font(c.add_run(caption.split(". ", 1)[0] + ". "), size=8.5, bold=True)
    if ". " in caption:
        set_run_font(c.add_run(caption.split(". ", 1)[1]), size=8.5, bold=False)


FIGURE_MAP = {
    "The general principle": (
        "figure_1_epistemic_trap",
        "Fig 1. Safety behaviour controls both risk and evidence. "
        "The baseline account holds contingency beliefs fixed, whereas the Bayesian extension closes the loop from posterior to policy, outcomes and learning. When the repetition count k never varies, baseline hazard q and preventive efficacy p cannot be separated."
    ),
    "Fixed action counts do not identify hazard and efficacy": (
        "figure_2_identifiability",
        "Fig 2. Structural non-identifiability under repetitive policies. "
        "(A) Each fixed repetition count defines an iso-likelihood ridge through the true parameters. (B) Every single-count design is singular, whereas policies with support on multiple counts are full rank. (C) Any two distinct counts identify both parameters; the diagonal remains singular."
    ),
    "A repetitive learner predicts outcomes without learning their cause": (
        "figure_3_learning",
        "Fig 3. A repetitive learner predicts outcomes without learning their cause. "
        "Posterior means and 10th-90th percentile bands are shown across 20 agents per condition over 2,000 episodes. A fixed policy leaves the posterior on a ridge and far from the true parameters (star); a variable policy recovers both efficacy and baseline hazard."
    ),
    "Informative contrasts restore causal knowledge": (
        "figure_4_interventions",
        "Fig 4. Interventions restore knowledge by creating informative contrasts. "
        "All panels show the same six conditions, identified redundantly by colour, line style and marker. "
        "(A) Efficacy-belief trajectories across baseline, intervention (shaded) and follow-up. "
        "(B) Structural knowledge gained versus harm experienced. (C) A post-action state probe can remove current-state uncertainty while adding little knowledge of the latent causes. "
        "(D) Change in repetition count at follow-up."
    ),
    "A matched-exposure experiment can test the account": (
        "figure_5_predictions_boundary",
        "Fig 5. A matched-exposure test in humans. "
        "(A) Four randomized action schedules have identical mean exposure; {1,5} and uniform 0-6 also have identical raw variance. "
        "(B,C) Simulated posterior uncertainty and efficacy error for those schedules. "
        "(D) Proposed measurement sequence and preregistered primary contrast."
    ),
}

SPECIAL_FIRST_PARAGRAPHS = {
    "Fixed action counts do not identify hazard and efficacy": (
        "Figure 2 visualises the structural result. Panel A shows that every fixed repetition count "
        "defines an iso-likelihood ridge through the true parameters rather than a unique peak. Panel B "
        "compares policies at 500 episodes: all four fixed-k policies are singular, whereas alternating "
        "between k in {0,1} yields log10 det I = 5.68 and graded variation over k in {0,...,6} yields 4.96. "
        "A high-count jittered policy over {5,6,7} is formally full rank but practically weak "
        "(log10 det I = -1.58), because harm becomes too rare to be informative. Panel C shows the same "
        "result over every two-point policy: the diagonal k1 = k2 is singular, and every off-diagonal "
        "contrast identifies both parameters. Variation matters, but only where outcomes still occur."
    ),
    "Informative contrasts restore causal knowledge": (
        "The simulation contains eight arms. Figure 4 displays the six intervention arms that share the "
        "same repetitive pre-treatment history and are therefore directly comparable. Continued free choice "
        "serves as the untreated reference in the reported contrasts; response prevention from a naive start "
        "is analysed separately because it lacks that history. The 11 pre-specified simulation contrasts and "
        "their operational outcomes are listed in results/phase3_hypotheses.csv."
    ),
}


def add_parameter_table(doc):
    caption = doc.add_paragraph("Table 1. Canonical model and simulation parameters", style="Caption")
    caption.paragraph_format.keep_with_next = True
    rows = [
        ("Parameter", "Default", "Interpretation"),
        ("q", "0.30", "Probability that an episode begins in danger"),
        ("p", "0.80", "Probability that one preventive repetition removes danger"),
        ("k", "0-10", "Number of preventive repetitions"),
        ("Grid", "101 x 101", "Exact posterior grid over (q, p)"),
        ("Prior: severe doubt", "Beta(1.5, 13.5)", "Initial mean efficacy E[p] = 0.10"),
        ("Harm / repetition cost", "20.0 / 0.25", "Subjective decision costs"),
        ("Choice precision", "12", "Inverse temperature of the softmax policy"),
        ("Primary horizon", "2,000 episodes", "Bayesian learning simulations"),
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, [1900, 1900, 5560])
    header_props = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    header_props.append(header)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 1 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run(value), size=8.4, bold=(i == 0), color="000000")
            if i == 0:
                set_cell_shading(cell, PALE_BLUE)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_accounts_table(doc):
    caption = doc.add_paragraph("Table 2. The present account in relation to neighbouring theories", style="Caption")
    caption.paragraph_format.keep_with_next = True
    rows = [
        ("Account", "Proposed locus", "Difference from the present account"),
        ("Belief bias", "False point belief about preventive efficacy", "Belief updating is exact; the policy makes the evidence non-identifying."),
        ("Transition uncertainty", "Uncertainty about state transitions", "The transition learner need not be impaired."),
        ("Eligibility-trace imbalance", "Biased credit assignment", "Credit assignment is exact."),
        ("Habitual avoidance", "Over-reliance on habits", "The information failure can precede habit formation."),
        ("Safety behaviour", "Blocked disconfirmation", "The blocking mechanism is formalised as rank deficiency in a specific parameter pair."),
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, [2050, 2850, 4460])
    header_props = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    header_props.append(header)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run(value), size=8.1, bold=(i == 0), color="000000")
            if i == 0:
                set_cell_shading(cell, PALE_BLUE)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_box(doc, label, text, entries, proof=False):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT if proof else PALE_BLUE)
    set_cell_margins(cell, top=190, start=220, bottom=190, end=220)
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    edges = ("left",) if proof else ("top", "left", "bottom", "right")
    for edge in edges:
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "14" if not proof else "20")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), MUTED if proof else BLUE)
        borders.append(border)
    tc_pr.append(borders)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    set_run_font(p.add_run(latex_plain(label, entries) + ". "), size=10.2, bold=True,
                 color=BLUE if not proof else MUTED)
    parts = re.split(r"\[\[EQ:(\d+)\]\]", text)
    for index, part in enumerate(parts):
        if index % 2:
            equation_p = cell.add_paragraph()
            equation_p.paragraph_format.space_before = Pt(3)
            equation_p.paragraph_format.space_after = Pt(4)
            add_equation_native(doc, int(part), paragraph=equation_p)
        elif part:
            target = p if index == 0 else cell.add_paragraph()
            target.paragraph_format.space_before = Pt(0)
            target.paragraph_format.space_after = Pt(0)
            target.paragraph_format.line_spacing = 1.15
            set_run_font(target.add_run(latex_plain(part, entries)), size=10.2, color=INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def decode_latex_text(text):
    replacements = {
        r"{\AA}": "Å", r"{\aa}": "å", r'{\"o}': "ö", r'{\"O}': "Ö",
        r'{\"u}': "ü", r'{\"U}': "Ü", r'{\"a}': "ä", r'{\"A}': "Ä",
        r"\&": "&",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text.replace("{", "").replace("}", "")


def format_reference(entry):
    author = decode_latex_text(entry.get("author", "Unknown author")).replace(" and others", " et al.")
    author = author.replace(" and ", "; ")
    year = entry.get("year", "n.d.")
    title = decode_latex_text(entry.get("title", "Untitled"))
    journal = decode_latex_text(entry.get("journal", entry.get("booktitle", "")))
    volume = entry.get("volume", "")
    number = entry.get("number", "")
    pages = entry.get("pages", "").replace("--", "-")
    doi = entry.get("doi", "")
    source = journal
    if volume:
        source += f", {volume}"
    if number:
        source += f"({number})"
    if pages:
        source += f", {pages}"
    text = f"{author} ({year}). {title}. {source}."
    if doi:
        text += f" https://doi.org/{doi}"
    return re.sub(r"\s+", " ", text).strip()


def add_references(doc, entries):
    doc.add_heading("References", level=1)
    for key, entry in sorted(entries.items(), key=lambda item: (
            item[1].get("author", "").lower(), item[1].get("year", ""))):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.05
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run_font(p.add_run(format_reference(entry)), size=8.5, color=INK)


def add_supplementary_figures(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    doc.add_heading("Supplementary Figures", level=1)
    supplements = [
        ("fig4_phase_diagram.png", "Supplementary Figure S1. Phase diagrams over initial doubt, harm cost and information-seeking weight. The realised repetition count reveals a U-shaped relation between action and self-knowledge."),
        ("fig5_hysteresis.png", "Supplementary Figure S2. Path dependence during gradual and abrupt removal of danger. Fixed policies are especially slow to recalibrate when the environment becomes safer."),
        ("fig_S2_uncertainty_diagram.png", "Supplementary Figure S3. Prior uncertainty, harm cost, residual efficacy error and behavioural intensity."),
    ]
    for index, (filename, caption) in enumerate(supplements):
        if index:
            doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(ROOT / "figures" / filename), width=Inches(7.0))
        set_picture_alt(doc, caption.split(". ", 1)[0], caption)
        c = doc.add_paragraph(style="Caption")
        set_run_font(c.add_run(caption.split(". ", 1)[0] + ". "), size=8.4, bold=True)
        set_run_font(c.add_run(caption.split(". ", 1)[1]), size=8.4, bold=False)


def build():
    all_entries = parse_bib(BIB.read_text(encoding="utf-8"))
    tex_source = TEX.read_text(encoding="utf-8")
    cited_keys = []
    for match in re.finditer(r"\\cite[tp]\{([^}]+)\}", tex_source):
        for key in (item.strip() for item in match.group(1).split(",")):
            if key in all_entries and key not in cited_keys:
                cited_keys.append(key)
    entries = {key: all_entries[key] for key in cited_keys}
    macros = parse_macros(NUMBERS.read_text(encoding="utf-8"))
    blocks = prepare_blocks(tex_source, macros)

    doc = Document()
    configure_styles(doc)
    add_title_page(doc)

    pending = None
    phead = None
    current_heading = None
    list_pending = False
    first_paragraph_in_heading = False
    parameter_table_added = False
    accounts_table_added = False
    for kind, value in blocks:
        if kind in {"h1", "h2", "h3"}:
            heading = latex_plain(value, entries)
            if heading == "Model":
                heading = "Methods and Model"
            level = {"h1": 1, "h2": 2, "h3": 3}[kind]
            doc.add_heading(heading, level=level)
            current_heading = heading
            pending = FIGURE_MAP.get(heading)
            phead = None
            first_paragraph_in_heading = True
            continue
        if kind == "phead":
            phead = latex_plain(value, entries)
            continue
        if kind == "equation":
            number = equation_number(value)
            add_equation_native(doc, number)
            continue
        if kind == "theorem":
            label, text = value
            add_box(doc, label, text, entries)
            continue
        if kind == "proof":
            add_box(doc, "Proof", value, entries, proof=True)
            continue
        if kind == "item":
            list_pending = True
            continue

        text = latex_plain(value, entries)
        if not text:
            continue
        if first_paragraph_in_heading and current_heading in SPECIAL_FIRST_PARAGRAPHS:
            text = SPECIAL_FIRST_PARAGRAPHS[current_heading]
        first_paragraph_in_heading = False
        p = doc.add_paragraph(style="List Number" if list_pending else None)
        list_pending = False
        if phead:
            set_run_font(p.add_run(phead + " "), size=10.5, bold=True, color=INK)
            phead = None
        set_run_font(p.add_run(text), size=10.5, color=INK)

        if (current_heading == "The Bayesian agent"
                and "Table 1 lists the canonical model and simulation parameters" in text
                and not parameter_table_added):
            add_parameter_table(doc)
            parameter_table_added = True
        if (current_heading == "What the model establishes"
                and "Table 2 places this narrower account" in text
                and not accounts_table_added):
            add_accounts_table(doc)
            accounts_table_added = True
        if pending:
            add_figure(doc, *pending)
            pending = None

    if not parameter_table_added or not accounts_table_added:
        raise ValueError("Both main tables must follow their first in-text citations")

    doc.add_heading("Declarations", level=1)
    declarations = [
        ("Ethics statement", "This study analyses mathematical models and simulated agents only; no human participants, animals or identifiable personal data were involved."),
        ("Data and code availability", f"All simulations are deterministic given a seed. The complete source code, result tables, figure-generation scripts and manuscript sources are publicly available at {REPOSITORY}."),
        ("Funding", "The author received no specific funding for this work."),
        ("Author contributions", "Dai Miyazaki: Conceptualization, Formal analysis, Investigation, Methodology, Software, Validation, Visualization, Writing - original draft, Writing - review & editing."),
        ("Competing interests", "The author declares no competing interests."),
    ]
    for label, value in declarations:
        p = doc.add_paragraph()
        set_run_font(p.add_run(label + ". "), size=10.2, bold=True)
        set_run_font(p.add_run(value), size=10.2)

    add_references(doc, entries)
    add_supplementary_figures(doc)

    properties = doc.core_properties
    properties.title = TITLE
    properties.subject = "Computational model of policy-induced non-identifiability in compulsive safety behaviour"
    properties.author = AUTHOR
    properties.keywords = KEYWORDS
    properties.comments = "Journal manuscript generated from version-controlled simulation outputs."

    path = OUTPUT / "compulsions_as_epistemic_traps_draft.docx"
    doc.save(path)
    print(path)


if __name__ == "__main__":
    build()
